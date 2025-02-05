import openpyxl
import time

# 1) Getting the excel file
# This path is representing the excel file
file = "C:/Users/DENNIS/Desktop/Automation_Excel_Files/W06-V02 Move and Resize.xlsx"

# 2) Getting a workbook
# This method will load the work book from the file
workbook = openpyxl.load_workbook(file)

# 3) Extracting a sheet from a work book
sheet = workbook["Sales 2016"]

# 4) Getting the number of rows and columns

# Returns the number of rows
rows = sheet.max_row
# Returns the number of columns
columns = sheet.max_column

# Reading all the rows and columns from excel sheet
for row in range(1, rows + 1):
    for column in range(1, columns + 1):
        # Capturing cell value
        print(sheet.cell(row, column).value, end='                        ')
    print()


print("Rows :",rows, "Columns :", columns)
