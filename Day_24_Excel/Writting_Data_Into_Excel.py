import openpyxl
import time

# SAME DATA
# -------------------------------------------------------------------------
# 1) Getting the excel file
# This path is representing the excel file
# file = "C:/Users/DENNIS/Desktop/Automation_Excel_Files/Book1.xlsx"
# workbook = openpyxl.load_workbook(file)
# # sheet=workbook.active["Sheet1"]
#
# # Here we needed not use the line above the reason being that in the file, there is only one sheet
# sheet=workbook.active
# for row in range(1,6):
#     for column in range(1,4):
#         # Writing to the excel sheet
#         sheet.cell(row, column).value="Welcome"
#         # print(sheet.cell(row, column).value)
#
#
# # This line makes sure that the data is saved
# workbook.save(file)

#-------------------------------------------------------------------------------

# DIFFERENT DATA
file = "C:/Users/DENNIS/Desktop/Automation_Excel_Files/Book2.xlsx"
workbook = openpyxl.load_workbook(file)
# sheet=workbook.active["Sheet1"]

# Here we needed not use the line above the reason being that in the file, there is only one sheet
sheet=workbook.active


sheet.cell(1,1).value=123 # First row first column
sheet.cell(1,2).value="Smith" # First row second column

sheet.cell(2,1).value=567
sheet.cell(2,2).value="John"

sheet.cell(3,1).value=345
sheet.cell(3,2).value="David"

sheet.cell(1,1).value=485
sheet.cell(1,2).value="Martha"
# This line makes sure that the data is saved
workbook.save(file)